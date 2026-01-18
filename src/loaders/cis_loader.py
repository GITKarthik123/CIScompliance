import openpyxl

def load_cis_controls(xlsx_path: str):
    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb.active

    headers = [c.value.strip() if c.value else "" for c in ws[1]]

    def col(name):
        if name not in headers:
            return None
        return headers.index(name)

    # 🔴 UPDATE THESE TO MATCH YOUR EXCEL EXACTLY
    IDX = {
        "Section": col("Domain"),
        "Control ID": col("CIS Control"),
        "Control Name": col("Control Title"),
        "Category": col("IG"),
        "Validation Type": col("Automation"),
    }

    controls = []

    for row in ws.iter_rows(min_row=2, values_only=True):
        controls.append({
            "Section": row[IDX["Section"]] if IDX["Section"] is not None else "",
            "Control ID": row[IDX["Control ID"]] if IDX["Control ID"] is not None else "",
            "Control Name": row[IDX["Control Name"]] if IDX["Control Name"] is not None else "",
            "Category": row[IDX["Category"]] if IDX["Category"] is not None else "",
            "Validation Type": (
                "AUTOMATED" if str(row[IDX["Validation Type"]]).upper() == "AUTOMATED"
                else "MANUAL"
            ) if IDX["Validation Type"] is not None else "MANUAL",
        })

    return controls
